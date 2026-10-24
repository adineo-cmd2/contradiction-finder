"""Ingest files (txt/md/pdf/docx) then run corpus-wide contradiction analysis."""

from io import BytesIO
from pathlib import Path

from sqlalchemy.orm import Session

from app.models.contradiction_record import ContradictionRecord
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.services.chunker import chunk_text
from app.services.contradiction_detector import detect
from app.services.contradiction_explainer import explain
from app.services.embedder import embed


def _decode_text(raw: bytes) -> str:
    return (raw or b"").decode("utf-8", errors="ignore")


def _extract_pdf(raw: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(raw or b""))
    texts = [(page.extract_text() or "") for page in reader.pages]
    return ("\n\n".join(t for t in texts if t.strip()), len(reader.pages))


def _extract_docx(raw: bytes) -> str:
    try:
        import docx  # type: ignore

        doc = docx.Document(BytesIO(raw or b""))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        pass
    # stdlib fallback: docx is a zip containing word/document.xml
    try:
        import xml.etree.ElementTree as ET
        import zipfile

        with zipfile.ZipFile(BytesIO(raw or b"")) as zf:
            xml_bytes = zf.read("word/document.xml")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        root = ET.fromstring(xml_bytes)
        paras = []
        for p in root.findall(".//w:p", ns):
            texts = [t.text or "" for t in p.findall(".//w:t", ns)]
            joined = "".join(texts).strip()
            if joined:
                paras.append(joined)
        return "\n\n".join(paras)
    except Exception:
        return ""


def ingest_file(
    db: Session, corpus_id: int, filename: str, raw_bytes: bytes, content_type: str = ""
) -> int:
    ext = Path(filename or "").suffix.lower()
    ctype = (content_type or "").lower()
    page_count = 0
    if ext == ".pdf" or ctype == "application/pdf":
        text, page_count = _extract_pdf(raw_bytes)
    elif ext == ".docx" or "wordprocessingml" in ctype:
        text = _extract_docx(raw_bytes)
    else:
        text = _decode_text(raw_bytes)

    doc = Document(corpus_id=corpus_id, filename=filename, page_count=page_count)
    db.add(doc)
    db.commit()
    db.refresh(doc)

    for idx, chunk in enumerate(chunk_text(text)):
        db.add(
            DocumentChunk(
                document_id=doc.id,
                idx=idx,
                text=chunk,
                embedding=embed(chunk),
                section=f"p{idx}",
            )
        )
    db.commit()
    return doc.id


def run_analysis(db: Session, corpus_id: int) -> int:
    docs = db.query(Document).filter(Document.corpus_id == corpus_id).all()
    doc_ids = [d.id for d in docs]
    db.query(ContradictionRecord).filter(ContradictionRecord.corpus_id == corpus_id).delete()
    db.commit()
    if not doc_ids:
        return 0
    chunks = (
        db.query(DocumentChunk).filter(DocumentChunk.document_id.in_(doc_ids)).all()
    )
    if not chunks:
        return 0
    payload = [
        {
            "id": c.id,
            "document_id": c.document_id,
            "text": c.text or "",
            "embedding": c.embedding or embed(c.text or ""),
        }
        for c in chunks
    ]
    detections = detect(payload)
    chunk_map = {c.id: c for c in chunks}
    for det in detections:
        ca = chunk_map[det["a_id"]]
        cb = chunk_map[det["b_id"]]
        info = explain(ca.text or "", cb.text or "", det["confidence"])
        db.add(
            ContradictionRecord(
                corpus_id=corpus_id,
                chunk_a_id=det["a_id"],
                chunk_b_id=det["b_id"],
                claim_a=info["claim_a"],
                claim_b=info["claim_b"],
                severity=info["severity"],
                explanation=info["explanation"],
                confidence=det["confidence"],
                same_doc=det["same_doc"],
            )
        )
    db.commit()
    return len(detections)
